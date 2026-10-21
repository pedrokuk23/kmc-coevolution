from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
wb=Workbook()
F="Arial"
H=Font(name=F,bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor="1F3A5F")
INP=Font(name=F,color="0000FF"); YEL=PatternFill("solid",fgColor="FFF2CC")
LNK=Font(name=F,color="008000"); NRM=Font(name=F); BLD=Font(name=F,bold=True)
TTL=Font(name=F,bold=True,size=14); NOTE=Font(name=F,italic=True,color="555555",size=9)
thin=Side(style="thin",color="BBBBBB"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
def hdr(ws,row,cols,start=1):
    for i,c in enumerate(cols):
        x=ws.cell(row=row,column=start+i,value=c); x.font=H; x.fill=HF; x.alignment=Alignment(wrap_text=True,vertical="center"); x.border=BOX
def inp(c,v):
    c.value=v; c.font=INP; c.fill=YEL; c.border=BOX
def frm(c,v,fmt=None,font=NRM):
    c.value=v; c.font=font; c.border=BOX
    if fmt: c.number_format=fmt
# ---------------- README
ws=wb.active; ws.title="README"
lines=[("KMC Worksheet — companion to Supplementary Material S1",TTL),
("Implements the calibration (Table 11) and the headroom procedure (Table 6) of the paper with spreadsheet formulas.",NRM),
("",NRM),("How to use",BLD),
("1. Calib_Hazard: enter the K assessed before each past integrated test, its cost and its outcome. The sheet fits the test hazard by likelihood and combines it with a public prior.",NRM),
("2. Calib_Rates: enter campaign data (cost vs K or M) and plateaus. The sheet fits the three rates and the two ceilings.",NRM),
("3. Parameters: calibrated values flow in automatically (green). Type a value in the Override column to use a rounded or prior value instead. Set the target, exposure limit and test cost.",NRM),
("4. Decision: at each review, enter cumulative cost, K and M. The sheet computes the ceilings, rates, hazards, net knowledge rates and the recommended next activity.",NRM),
("",NRM),("Colour code",BLD),("Blue text on yellow: inputs you edit.   Black: formulas.   Green: values linked from another sheet.",NRM),
("",NRM),("Example data",BLD),("All example data belong to the fictional organization 'Org X' of Supplementary Material S1 (Rocket R, engine E-1, Project N). Replace them with your own records.",NRM),
("",NRM),("Limits",BLD),("The sheet applies the observable procedure. The optimal policy, the exact derivative D_a and Monte Carlo checks need the Python code (kmc_model.py).",NRM)]
for i,(t,f) in enumerate(lines,1):
    c=ws.cell(row=i,column=1,value=t); c.font=f; c.alignment=Alignment(wrap_text=True,vertical="top")
ws.column_dimensions["A"].width=120
# ---------------- Calib_Hazard
ws=wb.create_sheet("Calib_Hazard")
ws["A1"]="Test hazard λ_T: likelihood of past integrated tests + public prior (S1, Sec. 3.1–3.2)"; ws["A1"].font=TTL
ws["A3"]="Prior median λ (per BU)"; inp(ws["B3"],0.90); ws["C3"]="FAA first flights 1980–2002 (paper, Sec. VI-B)"
ws["A4"]="Prior σ of ln λ"; inp(ws["B4"],0.5); ws["C4"]="assumed spread of the prior"
for r in (3,4): ws.cell(row=r,column=1).font=NRM; ws.cell(row=r,column=3).font=NOTE
hdr(ws,6,["Test n","K before test","Test cost c_f (BU)","Outcome (1 = failure/shock, 0 = success)","x = (1−K)²·c_f","failure flag","success flag","ΔK to next test","(1−K)·c_f","ΔK·(1−K)c_f","((1−K)c_f)²"])
Ks=[0.10,0.16,0.21,0.27,0.32,0.37,0.42,0.47]; outs=[1,1,1,1,1,1,1,0]
N=20
for i in range(N):
    r=7+i
    ws.cell(row=r,column=1,value=i+1).font=NRM
    if i<8:
        inp(ws.cell(row=r,column=2),Ks[i]); inp(ws.cell(row=r,column=3),2); inp(ws.cell(row=r,column=4),outs[i])
    else:
        for c in (2,3,4): inp(ws.cell(row=r,column=c),None)
    frm(ws.cell(row=r,column=5),f'=IF(B{r}="",0,(1-B{r})^2*C{r})',"0.000")
    frm(ws.cell(row=r,column=6),f'=IF(B{r}="",0,IF(D{r}=1,1,0))',"0")
    frm(ws.cell(row=r,column=7),f'=IF(B{r}="",0,IF(D{r}=0,1,0))',"0")
    if i<N-1:
        frm(ws.cell(row=r,column=8),f'=IF(OR(B{r}="",B{r+1}=""),0,B{r+1}-B{r})',"0.000")
        frm(ws.cell(row=r,column=9),f'=IF(OR(B{r}="",B{r+1}=""),0,(1-B{r})*C{r})',"0.000")
    else:
        frm(ws.cell(row=r,column=8),0,"0.000"); frm(ws.cell(row=r,column=9),0,"0.000")
    frm(ws.cell(row=r,column=10),f'=H{r}*I{r}',"0.0000"); frm(ws.cell(row=r,column=11),f'=I{r}^2',"0.0000")
last=7+N-1
# results
R=29
ws.cell(row=R,column=1,value="Results").font=BLD
res=[("λ_T, data only (maximum likelihood)",f"=INDEX(N$7:N$506,MATCH(MAX(O$7:O$506),O$7:O$506,0))","0.00"),
("λ_T, posterior median (use this)",f"=INDEX(N$7:N$506,MATCH(0.5,S$7:S$506,1)+1)","0.00"),
("posterior 5th percentile",f"=INDEX(N$7:N$506,MATCH(0.05,S$7:S$506,1)+1)","0.00"),
("posterior 95th percentile",f"=INDEX(N$7:N$506,MATCH(0.95,S$7:S$506,1)+1)","0.00"),
("e_T, knowledge per BU of integrated testing",f"=IF(SUM(K7:K{last})=0,\"\",SUM(J7:J{last})/SUM(K7:K{last}))","0.000")]
for i,(a,f,fm) in enumerate(res):
    ws.cell(row=R+1+i,column=1,value=a).font=NRM; frm(ws.cell(row=R+1+i,column=2),f,fm,BLD)
ws.cell(row=R+7,column=1,value="Grid search on columns N–S (λ from 0.05 to 5.04 per BU, step 0.01). Extend the grid if the estimate sits at its edge.").font=NOTE
hdr(ws,6,["λ grid","log-likelihood","log-prior","log-posterior (shifted)","posterior weight","cumulative"],start=14)
for i in range(500):
    r=7+i
    frm(ws.cell(row=r,column=14),0.05+0.01*i if i==0 else f"=N{r-1}+0.01","0.00")
    frm(ws.cell(row=r,column=15),f"=SUMPRODUCT($F$7:$F${last},LN(1-EXP(-N{r}*$E$7:$E${last})+1E-12))-N{r}*SUMPRODUCT($G$7:$G${last},$E$7:$E${last})","0.000")
    frm(ws.cell(row=r,column=16),f"=-0.5*((LN(N{r})-LN($B$3))/$B$4)^2-LN(N{r})","0.000")
    frm(ws.cell(row=r,column=17),f"=O{r}+P{r}","0.000")
    frm(ws.cell(row=r,column=18),f"=EXP(Q{r}-MAX($Q$7:$Q$506))","0.0000")
    frm(ws.cell(row=r,column=19),f"=SUM($R$7:R{r})/SUM($R$7:$R$506)","0.0000")
for col,w in zip("ABCDEFGHIJKNOPQRS",[38,12,12,16,12,9,9,12,12,12,12,9,12,10,12,11,11]): ws.column_dimensions[col].width=w
ws.row_dimensions[6].height=48
# ---------------- Calib_Rates
ws=wb.create_sheet("Calib_Rates")
ws["A1"]="Rates and ceilings (S1, Sec. 3.3–3.4)"; ws["A1"].font=TTL
def block(r0,title,xlab,ylab,xs,ys,ylog,rate_name,extra=None):
    ws.cell(row=r0,column=1,value=title).font=BLD
    hdr(ws,r0+1,[xlab,ylab,"ln(gap)"])
    for i in range(8):
        r=r0+2+i
        inp(ws.cell(row=r,column=1),xs[i] if i<len(xs) else None); inp(ws.cell(row=r,column=2),ys[i] if i<len(ys) else None)
        frm(ws.cell(row=r,column=3),f'=IF(B{r}="","",LN({ylog.format(r=r)}))',"0.000")
    a,b=r0+2,r0+9
    ws.cell(row=r0+2,column=5,value=rate_name).font=NRM
    frm(ws.cell(row=r0+2,column=6),f"=-SLOPE(C{a}:C{b},A{a}:A{b})","0.000",BLD)
    return a,b
ws["H3"]="Ceiling levels used by the S and T campaigns"; ws["H3"].font=BLD
ws["H4"]="M during the CFD campaign"; inp(ws["I4"],0.10)
ws["H5"]="K during the ground-test campaign"; inp(ws["I5"],0.50)
ws["H6"]="K̄ at that M (from the ceiling fit)"; frm(ws["I6"],"=F32+F33*I4","0.000")
ws["H7"]="M̄ at that K (from the ceiling fit)"; frm(ws["I7"],"=F42+F43*I5","0.000")
block(3,"Separate-effects campaign (engine E-1): K versus cost","cost (BU)","K",[0,1,2,3,4],[.30,.37,.43,.48,.53],"1-B{r}","r_E (per BU)")
ws["E9"]="(ln(1−K) = ln(1−K₀) − r_E·c)"; ws["E9"].font=NOTE
block(13,"Simulation campaign (CFD): K versus cost","cost (BU)","K",[0,1,2,3],[.20,.27,.32,.36],"$I$6-B{r}","r_S (per BU)")
ws["E19"]="(ln(K̄−K) = ln(K̄−K₀) − r_S·c)"; ws["E19"].font=NOTE
block(46,"Integrated campaign (ground tests): M versus cost","cost (BU)","M",[0,1,2,3],[.10,.20,.28,.34],"$I$7-B{r}","r_T (per BU)")
ws["E52"]="(ln(M̄−M) = ln(M̄−M₀) − r_T·c)"; ws["E52"].font=NOTE
# ceilings
ws["A24"]="Simulation plateaus: K where S stopped improving, versus demonstrated M"; ws["A24"].font=BLD
hdr(ws,25,["M","plateau K"])
for i,(m,k) in enumerate(zip([.10,.30,.50,.70],[.47,.57,.71,.82])):
    inp(ws.cell(row=26+i,column=1),m); inp(ws.cell(row=26+i,column=2),k)
for i in range(4,6): inp(ws.cell(row=26+i,column=1),None); inp(ws.cell(row=26+i,column=2),None)
ws["E32"]="k_b (intercept)"; frm(ws["F32"],"=INTERCEPT(B26:B31,A26:A31)","0.000",BLD)
ws["E33"]="k_l (slope)"; frm(ws["F33"],"=SLOPE(B26:B31,A26:A31)","0.000",BLD)
ws["A34"]="Test plateaus: M where T stopped passing more envelope, versus K"; ws["A34"].font=BLD
hdr(ws,35,["K","plateau M"])
for i,(k,m) in enumerate(zip([.30,.45,.60,.75],[.36,.53,.67,.84])):
    inp(ws.cell(row=36+i,column=1),k); inp(ws.cell(row=36+i,column=2),m)
for i in range(4,6): inp(ws.cell(row=36+i,column=1),None); inp(ws.cell(row=36+i,column=2),None)
ws["E42"]="m_b (intercept)"; frm(ws["F42"],"=INTERCEPT(B36:B41,A36:A41)","0.000",BLD)
ws["E43"]="m_l (slope)"; frm(ws["F43"],"=SLOPE(B36:B41,A36:A41)","0.000",BLD)
for col,w in zip("ABCDEFGHI",[16,14,12,3,34,12,3,36,10]): ws.column_dimensions[col].width=w
# ---------------- Parameters
ws=wb.create_sheet("Parameters")
ws["A1"]="Parameter card (paper Table 5 notation)"; ws["A1"].font=TTL
hdr(ws,3,["Symbol","Meaning","Calibrated / default","Override (optional)","Value used"])
P=[("k_b","simulation ceiling, intercept","=Calib_Rates!F32",0.40),("k_l","simulation ceiling, slope","=Calib_Rates!F33",0.60),
("m_b","test ceiling, intercept","=Calib_Rates!F42",0.05),("m_l","test ceiling, slope","=Calib_Rates!F43",1.05),
("r_S","simulation rate (per BU)","=Calib_Rates!F15",0.30),("r_E","separate-effects rate (per BU)","=Calib_Rates!F5",0.10),
("e_E","separate-effects maturity credit (per BU)",0.05,None),("r_T","integrated-test maturity rate (per BU)","=Calib_Rates!F48",0.24),
("e_T","integrated-test knowledge credit (per BU)","=Calib_Hazard!B34",0.035),("λ_E","separate-effects shock hazard scale (per BU)",0.05,None),
("λ_T","integrated-test shock hazard scale (per BU)","=Calib_Hazard!B31",1.5),("δ_K","K lost in a shock",0.06,None),("δ_M","M lost in a shock",0.15,None),("ρ","rework per shock (BU)",2.0,None)]
for i,(s,m,cal,ov) in enumerate(P):
    r=4+i
    ws.cell(row=r,column=1,value=s).font=BLD; ws.cell(row=r,column=2,value=m).font=NRM
    c=ws.cell(row=r,column=3)
    if isinstance(cal,str): frm(c,cal,"0.000",LNK)
    else: inp(c,cal); c.number_format="0.000"
    inp(ws.cell(row=r,column=4),ov); ws.cell(row=r,column=4).number_format="0.000"
    frm(ws.cell(row=r,column=5),f'=IF(D{r}="",C{r},D{r})',"0.000",BLD)
ws["D4"].comment=Comment("Override column: Org X rounds the calibrated values to the card of S1, Sec. 3.5. Clear a cell to use the calibrated value.","KMC")
ws["A19"]="Project and policy"; ws["A19"].font=BLD
proj=[("K_req","required knowledge (target region)",0.75),("M_req","required maturity (target region)",0.80),("h_max","exposure limit: λ_T(1−K)² per BU",0.05),("c_f","cost of one integrated test (BU)",2.0),("Budget","program budget (BU)",30)]
for i,(s,m,v) in enumerate(proj):
    r=20+i; ws.cell(row=r,column=1,value=s).font=BLD; ws.cell(row=r,column=2,value=m).font=NRM; inp(ws.cell(row=r,column=5),v)
ws["A26"]="Thresholds (S1, Sec. 4.2)"; ws["A26"].font=BLD
thr=[("K_need","knowledge for which M̄(K) reaches M_req","=MIN(1,MAX(0,(E21-E6)/E7))"),("K_haz","knowledge for which the test hazard is within h_max","=MIN(1,MAX(0,1-SQRT(E22/E14)))"),("K_test","test threshold = max(K_need, K_haz)","=MAX(E27,E28)"),("P(lose one test) at K_test","1 − exp(−h_max·c_f) at the threshold","=1-EXP(-E14*(1-E29)^2*E23)")]
for i,(s,m,f) in enumerate(thr):
    r=27+i; ws.cell(row=r,column=1,value=s).font=BLD; ws.cell(row=r,column=2,value=m).font=NRM; frm(ws.cell(row=r,column=5),f,"0.000",BLD)
ws["E30"].number_format="0.0%"
for col,w in zip("ABCDE",[26,48,20,20,14]): ws.column_dimensions[col].width=w
# ---------------- Decision
ws=wb.create_sheet("Decision")
ws["A1"]="Decision at each review — headroom procedure (Table 6, steps 3–4)"; ws["A1"].font=TTL
ws["A2"]="Enter date, cumulative cost, K and M after each review (yellow). Example rows: Project N of S1, Sec. 4.4."; ws["A2"].font=NOTE
cols=["Review","Cum. cost (BU)","K","M","K̄(M)","M̄(K)","S: K gain","E: K gain","T: K gain","E: hazard","T: hazard","net S","net E","net T","Step","Recommended activity","P(lose a test now)","Budget left (BU)"]
hdr(ws,4,cols)
ex=[("R0",0,.200,.100),("R1",4,.465,.156),("R2",8,.642,.244),("R3",12,.760,.343),("R4",14.8,.819,.414),("R5",18,.838,.688),("R6",20.4,.85,.80)]
Pm="Parameters!$E$"
kb,kl,mb,ml,rS,rE,eE,rT,eT,lE,lT,dK=[Pm+str(4+i) for i in range(12)]
Kreq,Mreq,hmax,cf,bud=[Pm+str(20+i) for i in range(5)]; Kneed=Pm+"27"
for i in range(30):
    r=5+i
    if i<len(ex):
        for j,v in enumerate(ex[i]): inp(ws.cell(row=r,column=1+j),v)
    else:
        for j in range(4): inp(ws.cell(row=r,column=1+j),None)
    e=f'IF(C{r}="","",'
    frm(ws.cell(row=r,column=5),f'={e}MIN(1,{kb}+{kl}*D{r}))',"0.000")
    frm(ws.cell(row=r,column=6),f'={e}MIN(1,{mb}+{ml}*C{r}))',"0.000")
    frm(ws.cell(row=r,column=7),f'={e}{rS}*MAX(E{r}-C{r},0))',"0.0000")
    frm(ws.cell(row=r,column=8),f'={e}{rE}*(1-C{r}))',"0.0000")
    frm(ws.cell(row=r,column=9),f'={e}{eT}*(1-C{r}))',"0.0000")
    frm(ws.cell(row=r,column=10),f'={e}{lE}*(1-C{r})^2)',"0.000")
    frm(ws.cell(row=r,column=11),f'={e}{lT}*(1-C{r})^2)',"0.000")
    frm(ws.cell(row=r,column=12),f'={e}G{r})',"0.0000")
    frm(ws.cell(row=r,column=13),f'={e}H{r}-J{r}*{dK})',"0.0000")
    frm(ws.cell(row=r,column=14),f'={e}I{r}-K{r}*{dK})',"0.0000")
    frm(ws.cell(row=r,column=15),f'={e}IF(AND(C{r}>={Kreq},D{r}>={Mreq}),"done",IF(AND(D{r}<{Mreq},C{r}>={Kneed},K{r}<={hmax}),"3",IF(AND(C{r}>={Kreq},C{r}>={Kneed},K{r}<={hmax}),"3",IF(AND(C{r}>={Kneed},K{r}>{hmax}),"3b/4","4")))))')
    frm(ws.cell(row=r,column=16),f'={e}IF(O{r}="done","Target region reached: decision checkpoint",IF(O{r}="3","T  integrated test",IF(MAX(L{r}:N{r})=L{r},"S  simulation / analysis",IF(MAX(L{r}:N{r})=M{r},"E  separate-effects test","T  integrated test (to learn)")))))',None,BLD)
    frm(ws.cell(row=r,column=17),f'={e}1-EXP(-K{r}*{cf}))',"0%")
    frm(ws.cell(row=r,column=18),f'=IF(B{r}="","",{bud}-B{r})',"0.0")
ws["A37"]="Step 3: M < M_req, K ≥ K_need and hazard ≤ h_max → integrated test.  Step 4 (and 3b): otherwise the activity with the largest net K rate = K gain − hazard·δ_K.  Always screen the activity against cost, schedule and exposure (step 6) and re-assess after every result."
ws["A37"].font=NOTE
for col,w in zip("ABCDEFGHIJKLMNOPQR",[8,10,8,8,8,8,9,9,9,9,9,9,9,9,7,40,10,10]): ws.column_dimensions[col].width=w
ws.row_dimensions[4].height=34; ws.freeze_panes="E5"
wb.save("KMC_Worksheet.xlsx"); print("saved")
