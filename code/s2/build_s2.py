"""Builds Supplementary Material S2 (documented programs) as .docx; figures from figs_s2.py.
Run from this folder: python3 figs_s2.py && python3 build_s2.py [output.docx]. Convert to PDF with LibreOffice:
soffice --headless --convert-to pdf KMC_Supplementary_S2_Documented_Programs.docx"""
REFS = {
    14: "X-43A Mishap Investigation Board, \"Report of findings: X-43A mishap, vol. I,\" NASA, Washington, DC, USA, May 8, 2003. Accessed: Oct. 6, 2026. [Online]. Available: https://www.nasa.gov/wp-content/uploads/2015/01/47414main_x43A_mishap.pdf",
    15: "J. C. Oefelein and V. Yang, \"Comprehensive review of liquid-propellant combustion instabilities in F-1 engines,\" J. Propuls. Power, vol. 9, no. 5, pp. 657–677, Sep. 1993, doi: 10.2514/3.23674.",
    16: "T. Malik, \"SpaceX successfully launches Falcon 1 rocket into orbit,\" Space.com, Sep. 29, 2008. Accessed: Oct. 6, 2026. [Online]. Available: https://www.space.com/5905-spacex-successfully-launches-falcon-1-rocket-orbit.html",
    17: "J. Uri, \"55 years ago: NASA administrator Webb OKs flying a crew on the third Saturn V,\" NASA History, Apr. 27, 2023. Accessed: Oct. 6, 2026. [Online]. Available: https://www.nasa.gov/history/55-years-ago-nasa-administrator-webb-oks-flying-a-crew-on-the-third-saturn-v/",
    54: "V. L. Rausch, C. R. McClinton, J. Sitz, and P. Reukauf, \"NASA's Hyper-X program,\" presented at the 51st Int. Astronaut. Congr., Rio de Janeiro, Brazil, Oct. 2–6, 2000, Paper IAF-00-V.4.01. Accessed: Oct. 6, 2026. [Online]. Available: https://ntrs.nasa.gov/citations/20010075170",
    55: "D. E. Reubush, L. T. Nguyen, and V. L. Rausch, \"Review of X-43A return to flight activities and current status,\" presented at the 12th AIAA Int. Space Planes Hypersonic Syst. Technol. Conf., Norfolk, VA, USA, Dec. 2003, AIAA Paper 2003-7085. Accessed: Oct. 6, 2026. [Online]. Available: https://ntrs.nasa.gov/citations/20040040296",
    56: "NASA Dryden Flight Research Center, \"NASA's X-43A hypersonic aircraft ready for flight,\" Edwards, CA, USA, Release 04-14, Mar. 24, 2004. Accessed: Oct. 6, 2026. [Online]. Available: https://www.nasa.gov/news-release/nasa-dryden-flight-research-center-news-room-news-releases-nasas-x-43a-hypersonic-aircraft-ready-for-flight/",
    57: "NASA Dryden Flight Research Center, \"NASA's X-43A flight results in treasure trove of data,\" Edwards, CA, USA, Release 04-20, Apr. 7, 2004. Accessed: Oct. 6, 2026. [Online]. Available: https://www.nasa.gov/news-release/nasa-dryden-flight-research-center-news-room-news-releases-nasas-x-43a-flight-results-in-treasure-trove-of-data/",
    58: "C. Bahm, E. Baumann, J. Martin, D. Bose, R. E. Beck, and B. Strovers, \"The X-43A Hyper-X Mach 7 flight 2 guidance, navigation, and control overview and flight test results,\" presented at the AIAA/CIRA 13th Int. Space Planes Hypersonic Syst. Technol. Conf., Capua, Italy, May 2005. Accessed: Oct. 6, 2026. [Online]. Available: https://ntrs.nasa.gov/citations/20050182778",
    59: "L. Grindle and C. Bahm, \"The X-43A (Hyper-X) flies into the record books,\" in 2004 Research Engineering Annual Report, NASA Dryden Flight Research Center, Edwards, CA, USA, Jun. 2006. Accessed: Oct. 6, 2026. [Online]. Available: https://ntrs.nasa.gov/citations/20060056098",
    60: "NASA, \"X-43A Hyper-X,\" last updated Aug. 15, 2025. Accessed: Oct. 6, 2026. [Online]. Available: https://www.nasa.gov/reference/x-43a/",
    61: "NASA Dryden Flight Research Center, \"X-43A flight delayed,\" Edwards, CA, USA, Release 04-06, Feb. 17, 2004. Accessed: Oct. 6, 2026. [Online]. Available: https://www.nasa.gov/news-release/nasa-dryden-flight-research-center-news-room-news-releases-x-43a-flight-delayed/",
    62: "C. N. Dennehy, \"GN&C lessons learned and associated best practices,\" presented at ESA/ESTEC, Noordwijk, The Netherlands, Jun. 6, 2017. Accessed: Oct. 6, 2026. [Online]. Available: https://ntrs.nasa.gov/citations/20200009546",
    63: "G. E. Mueller, \"Revised manned space flight schedule,\" NASA memorandum M-C M 9330.186, Oct. 31, 1963. Accessed: Oct. 6, 2026. [Online]. Available: https://www.american-spacecraft.org/documents/sp-4503/document-7.html",
    64: "S. Clark, \"Falcon 1 suffers another setback,\" Spaceflight Now, Aug. 3, 2008. Accessed: Oct. 6, 2026. [Online]. Available: https://spaceflightnow.com/falcon/003/failure.html",
    65: "Air Force Research Laboratory, \"X-51A flight ends prematurely,\" Wright-Patterson AFB, OH, USA, Aug. 15, 2012. Accessed: Oct. 6, 2026. [Online]. Available: https://www.edwards.af.mil/News/AFMC-News/Article/153597/x-51a-flight-ends-prematurely/",
    66: "95th Air Base Wing Public Affairs, \"Second hypersonic flight ends prematurely, brings new flight test data,\" U.S. Air Force, Jun. 15, 2011. Accessed: Oct. 6, 2026. [Online]. Available: https://www.af.mil/News/Article-Display/Article/113041/second-hypersonic-flight-ends-prematurely-brings-new-flight-test-data/",
    67: "Boeing, \"Boeing X-51A WaveRider sets record with successful 4th flight,\" news release, May 3, 2013. Accessed: Oct. 6, 2026. [Online]. Available: https://boeing.mediaroom.com/2013-05-03-Boeing-X-51A-WaveRider-Sets-Record-with-Successful-4th-Flight",
    68: "Defense Advanced Research Projects Agency, \"14-F-0122 DARPA X-51 documents,\" FOIA release, Washington Headquarters Services FOIA Reading Room, 2014. Accessed: Oct. 6, 2026. [Online]. Available: https://www.esd.whs.mil/Portals/54/Documents/FOID/Reading%20Room/DARPA/14-F-0122_DARPA_X-51_DOCUMENTS.pdf",
    69: "Defense Advanced Research Projects Agency, \"DARPA HTV-2 second test flight report released,\" SpaceRef, Apr. 20, 2012. Accessed: Oct. 6, 2026. [Online]. Available: https://spaceref.com/newspace-and-tech/darpa-htv-2-second-test-flight-report-released/",
}

TITLE = "Supplementary Material S2: Documented Programs Read Through KMC"
SUBTITLE = ("Supplement to: P. K. de Albuquerque, D. R. Pleffken, C. S. Cerqueira, and L. Galembeck, "
            "\"Knowledge–Maturity Coevolution: A Coupled Evidence Model for Selecting Simulation and Test Activities "
            "in Constrained Engineering Development,\" IEEE Access.")

NOTE = ("This supplement expands Section VI-C of the article. Numbers in brackets refer to the article's reference list; "
        "the cited items are reproduced at the end. Section, table, and equation numbers without the prefix S2 refer to the article. "
        "All sources are public. The readings are retrospective and were made with the outcomes known; they illustrate whether the "
        "mechanisms of the model appear in documented decisions and do not calibrate the model or show that KMC would have changed "
        "any decision.")

# (kind, content) blocks; kinds: H1, H2, P, FIG, TABLE
BODY = [
("H1", "1. Selection and coding"),
("P", "Six aerospace programs with public, primary or near-primary documentation of the sequence of analyses, tests, "
      "discrepancies, and corrective actions were selected purposively: the Hyper-X Launch Vehicle (HXLV) of the X-43A program, "
      "the Saturn V all-up test decision, the F-1 combustion-instability program, Falcon 1, and the X-51A and DARPA Hypersonic "
      "Technology Vehicle 2 (HTV-2) programs. The selection is illustrative, not a representative sample of development programs."),
("P", "For each program, documented events were coded by activity class (S: analysis and simulation; E: separate-effects or "
      "component testing; T: integrated ground or flight testing), by whether a test produced an epistemic shock (a negative "
      "reassessment after evidence undermined a prior claim, Section III-B), and by the evidence available before the next "
      "integrated test. No numerical K or M was assigned, because the public record does not contain criterion-level evidence "
      "for complete profiles (see 2.2 below). Where the record names a specific deficiency, it is mapped to the criteria of the "
      "operational rubric (Table 4): for K, causal coverage, mechanistic depth, coherence, uncertainty understanding, and domain "
      "definition; for M, demonstrated performance, representativeness, integration fidelity, decision robustness, and control and "
      "repeatability."),

("H1", "2. X-43A Hyper-X Launch Vehicle (HXLV)"),
("H2", "2.1 Claims and documentary corpus"),
("P", "The case concerns the HXLV booster, not the X-43A research vehicle (HXRV); findings on one are not transferred to the other. "
      "Three claims were defined: C1, the aerodynamic and actuator models used for stability assessment; C2, the uncertainty "
      "procedure for the intended trajectory; and C3, boost control capability. C1 and C2 concern models and procedures, C3 the "
      "vehicle, and their maturity is not aggregated."),
("P", "The bounded corpus comprises eight public documents dated 2000–2025: the preflight program paper [54], the mishap "
      "investigation [14], the return-to-flight review [55], two NASA releases on the second flight [56], [57], the flight-2 "
      "guidance and control paper [58], the 2004 research engineering report [59], and the NASA historical page [60]. From them, "
      "19 parent records yield 20 observations, one of which (a conflicting release altitude) is retained only for discrepancy "
      "auditing. A 2004 incident release [61] and a lessons-learned review [62] were added later as a separately logged extension "
      "and did not change the original records."),
("P", "Each document entered the epochs by its publication date: [54] in E1, [14] in E2, [55] and [56] in E3, and [57]–[59] in "
      "E4. The NASA historical page [60], updated in 2025, belongs to no epoch: it is used only to record a conflicting release "
      "altitude and never enters an assessment, so no information published after June 2006 enters any epoch."),
("P", "Records were grouped in four cumulative documentary epochs: E1, preflight planning (October 2000); E2, investigation "
      "(through May 2003); E3, corrective actions (through March 2004); and E4, retrospective results (through June 2006). The "
      "epochs contain 3, 7, 14, and 19 observations; these cumulative counts organize the documents and are not independent "
      "samples. Publication dates delimit the epochs; they do not establish what individual "
      "engineers knew, and later reconstructions were excluded from earlier epochs. Repeated accounts of the same event were linked "
      "rather than counted as independent demonstrations. Conflicting values were preserved rather than averaged: the first-flight "
      "release altitude appears as about 20,000 ft [55], 24,000 ft [56], and 23,000 ft [60], and the 22,244 ft reported at "
      "failure time 13.5 s [14] is not a release altitude."),
("H2", "2.2 Summary of the documentary audit"),
("P", "The audit enumerates 120 criterion-level units, three claims × ten rubric criteria × four epochs, that is, twelve "
      "claim–epoch profiles of ten criteria each. For each unit it recorded the candidate passages, the configuration to which they "
      "apply, the missing information, and whether the evidence was sufficient to place the criterion on an anchor of Table 4. "
      "None was, so every unit remained not assessed and no profile, and hence no K or M, could be formed. Typical gaps were final model versions and "
      "complete parameter distributions (C1, C2), closure of reviews and acceptance records, repeatability, and activity-level costs "
      "(C3), and export-controlled material omitted from the public investigation [14]. The result supports the restriction of the "
      "article to qualitative readings; it does not show that the program lacked the evidence. This section summarizes the audit; "
      "the unit-level working records are not reproduced in this supplement."),
("H2", "2.3 Sequence"),
("P", "The preflight paper reported control-law refinement and Monte Carlo analysis of aerodynamic and flight-control "
      "uncertainties, and a trajectory different from that of the heritage booster [54]. After the loss of control during boost on "
      "2 June 2001, the investigation found inaccuracies in inherited and HXLV-specific models: corrected nominal models reduced the "
      "predicted roll-axis gain margin from 8 dB to less than 2 dB against a 6 dB requirement, and reproducing the mishap required "
      "combining the model corrections with parameter uncertainty [14] (Fig. S2-1). Return to flight combined wind-tunnel and "
      "actuator tests, revised models, comparisons of independent boost simulations, stronger actuators, and propellant off-loading "
      "for a higher-altitude release that reduced exposure [55], [56]. A test-setup anomaly on 11 February 2004 exceeded the "
      "qualified actuator torque and led to replacement and investigation before the second flight [61]. On 27 March 2004 the HXLV "
      "delivered the research vehicle at 28.67 km (94,069 ft) against 28.63 km (93,932 ft) planned and at Mach 6.946 against "
      "7.075 [57], [58]; the 2004 report states that both the Mach 7 and the Mach 10 flights delivered the research vehicle "
      "within the separation tolerances [59]. A lessons-learned "
      "review relates the accident to heritage assumptions, representative testing, stability margins, and uncertainty and dynamic "
      "modeling [62]. Fig. S2-2 separates the event dates from the dates of the documents that report them."),
("FIG", ("S2_Fig1.png", "Fig. S2-1. Reported X-43A roll-axis gain margin before and after nominal model correction [14, pp. 6–7]. "
         "The corrected value is reported only as less than 2 dB (open marker: upper bound); the requirement is 6 dB. These are "
         "model predictions, not flight measurements or KMC scores.")),
("FIG", ("S2_Fig2.png", "Fig. S2-2. Event dates and dates of selected documents for the first-flight mishap [14], the actuator "
         "incident [61], and the return-to-flight mission [57]. Connecting segments associate an event with a record; their "
         "lengths do not measure decision delay.")),
("H2", "2.4 Reading through the model"),
("P", "The margins of Fig. S2-1 are engineering quantities: 8 dB, 6 dB, and 2 dB correspond to gain factors of about 2.5, 2.0, "
      "and 1.26, and the corrected value is only an upper bound. They are not values of K or M, and the record gives no calibrated "
      "mapping between them; terms such as headroom or K_need below are readings through the framework, not reconstructed states. "
      "With that caveat, an integrated test was attempted while knowledge of the configuration-specific claims rested "
      "on heritage transfer, and the result can be read as an epistemic shock on C1–C3. Recovery came from the activities the model favors "
      "after a shock: separate-effects testing, revised models, and joint-uncertainty simulation, together with reduced exposure. "
      "Table S2-1 links the main observations to rubric criteria. Table S2-2 lists the postflight findings examined in the "
      "extension, and Table S2-3 a time-bounded check that uses only documents available before each flight, with cutoffs at the "
      "end of 1 June 2001 and 26 March 2004. The preflight sources support requests for evidence about domain transfer and "
      "uncertainty adequacy, but they do not establish that the model error was detectable before flight: the program, not KMC, "
      "detected and corrected the problems, and an equally informed conventional review could reach the same questions."),
("TABLE", ("Table S2-1. X-43A observations and KMC reasoning (retrospective).",
           ["Observation", "Claim and rubric criteria", "Implied action", "Inference excluded"],
           [["Corrected roll margin below requirement [14]", "C1 domain definition, coherence; C2 uncertainty understanding",
             "Reconcile model and uncertainty bounds before renewed acceptance", "No numerical drop in K or M; no claim of earlier prediction"],
            ["Nominal corrections alone did not reproduce the mishap [14]", "C2 decision robustness",
             "Require justified joint variations before closing the gate", "Uncertainty analysis already existed [54]"],
            ["Stronger actuators and modified trajectory [55], [56]", "C3 representativeness",
             "Declare the new baseline; verify changed interfaces and envelope", "No automatic recovery from completing a corrective action"],
            ["Actuator torque exceedance before flight 2 [61]", "C3 demonstrated performance",
             "Keep acceptance open until the replacement is verified", "Program detected it without KMC"],
            ["Flight-2 delivery close to target (94,069 vs 93,932 ft; Mach 6.946 vs 7.075) and reported within separation tolerances [58], [59]", "C3 demonstrated performance",
             "Credit the delivery point, not full-envelope repeatability", "No transfer of HXRV results or of the Mach 10 flight as a Mach 7 repetition"]],
           [1.6, 1.5, 1.7, 1.7])),
("TABLE", ("Table S2-2. Postflight findings and KMC reasoning (extension PF01–PF05).",
           ["Entry and evidence", "Rubric criteria", "Candidate response", "Inference limit"],
           [["PF01: inherited and HXLV-specific model deficiencies; corrected roll margin below requirement [14, pp. 6–7], [62]",
             "C1 domain definition, coherence; C3 representativeness", "Reopen model acceptance; require configuration-specific corroboration",
             "Post-mishap diagnosis; no demonstrated preflight signal"],
            ["PF02: corrected nominal models alone did not reproduce the mishap; combined uncertainty was required [14, pp. 6–7]",
             "C2 uncertainty understanding, decision robustness", "Require justified joint variations before closing the gate",
             "Uncertainty analysis already existed [54]; KMC does not discover omitted physics"],
            ["PF03: return-to-flight actions revised models, actuators, and launch conditions [55]",
             "C1–C3 separately: domain definition, representativeness", "Rebaseline affected claims; verify the changed configuration",
             "A design correction does not itself demonstrate recovered maturity"],
            ["PF04: 11 February 2004 test-setup anomaly exceeded qualified torque; replacement and investigation required [61]",
             "C3 demonstrated performance, representativeness", "Keep actuator acceptance open pending disposition and verification",
             "Program detected the anomaly without KMC"],
            ["PF05: flight-2 HXLV separation results and HXRV response discrepancies concern distinct objects [58]",
             "C3 demonstrated performance", "Accept only the supported delivery-point claim; keep out-of-scope discrepancies separate",
             "Success is not full-envelope repeatability; HXRV discrepancies cannot lower HXLV assessments"]],
           [2.0, 1.4, 1.6, 1.5])),
("TABLE", ("Table S2-3. Time-bounded documentary check before flights 1 and 2. G: evidence request; R: review trigger already "
           "recognized by the program; U: no eligible preflight indication.",
           ["ID / code", "Preflight record or exclusion", "KMC review implication", "What is established"],
           [["PR01 / G", "Flight 1: trajectory different from the heritage booster [54, pp. 2–3]", "Request evidence for transferring heritage models", "Domain difference known; specific error not established"],
            ["PR02 / G", "Flight 1: uncertainty analysis already reported [54, p. 3]", "Review bounds and combinations for the intended claim", "Existing analysis; inadequacy not established"],
            ["PR03 / G", "Flight 1: tests planned in the 2000 account [54, p. 5]", "Request dated completion and disposition records", "No evidence that tests remained unfinished at launch"],
            ["PR04 / U", "Flight 1: corrected margin comes from the 2003 report [14, pp. 6–7]", "Exclude hindsight from preflight diagnosis", "No eligible preflight indication"],
            ["PR05 / R", "Flight 2: actuator qualification exceeded [61, paras. 1–2]", "Require replacement and verification before closing acceptance", "Program already detected and investigated"],
            ["PR06 / R", "Flight 2: autopilot retuning considered [61, para. 3]", "Link revised parameters to verification evidence", "Program already recognized the review need"],
            ["PR07 / R", "Flight 2: corrective models, hardware, and simulations [55, pp. 4–7]", "Trace the changed configuration to evidence for each claim", "Corrective work already documented; no new defect inferred"]],
           [0.8, 2.1, 1.9, 1.7])),
("H2", "2.5 From evidence to the next activity"),
("P", "Two chains connect a documented observation to the choice of the next evidence-producing activity (Table S2-4). Chain A "
      "uses the investigation [14, pp. 6–7]: corrected nominal models remain stable and do not reproduce the loss of control, so "
      "coherence (C1) and uncertainty adequacy (C2) remain open, and the retained sequence couples targeted tests with "
      "joint-uncertainty simulation rather than closing on nominal stability. Chain B uses only the release of 17 February 2004 "
      "[61], before flight 2: exceeding the qualified torque removes the basis for accepting that actuator, so the retained "
      "sequence replaces and verifies it. For flight-1 boost control (C3), the record shows that the mandatory "
      "demonstrated-performance condition was not met, because control was lost. This is a failed demonstration, not the absence "
      "of one, and no anchor of Table 4 is assigned; consistently with the audit of 2.2, the criteria remain not assessed, so no "
      "(K, M) point or slope dM/dK can be computed; the criterion-based "
      "branch of the procedure remains usable where the derivative-based branch is not. In both chains the program reached the "
      "same actions without KMC; the alternatives are analytical comparisons, not a recovered historical trade study, and their "
      "costs and durations are unknown."),
("TABLE", ("Table S2-4. Candidate activities after the two documented observations. \"Retained\" denotes a plan for obtaining "
           "evidence; acceptance remains open until its completion criteria are verified.",
           ["ID", "Candidate activity", "Criterion-based disposition", "Evidence needed before closure"],
           [["A1", "Correct nominal models only", "Insufficient: does not reconcile the observed instability", "Explanation of the residual discrepancy"],
            ["A2", "Joint-uncertainty simulation alone", "Useful but conditional: unsupported bounds cannot establish adequacy", "Traceable bounds, model versions, and discriminating comparisons"],
            ["A3", "Targeted tests plus coupled model and uncertainty analysis", "Retained", "Measured support for key dynamics; traceable joint variations; applicability limits"],
            ["B1", "Keep the affected actuator because it may still function", "Excluded by the documented replacement requirement [61]", "Functioning alone does not dispose of the exceedance"],
            ["B2", "Retune the autopilot or simulate without replacing the unit", "Insufficient for the hardware qualification condition", "Disposition of the installed hardware"],
            ["B3", "Replace and verify the actuator; coordinate retuning and investigation", "Retained, consistent with the announced parallel activities [61]", "Replacement identity, acceptance results, and integrated control verification"]],
           [0.4, 2.0, 2.1, 2.0])),

("H1", "3. Saturn V all-up testing"),
("P", "On 31 October 1963, Mueller directed that the first Saturn IB and Saturn V flights use all live stages and complete "
      "spacecraft, and that mission planning assume two successful flights before a crewed flight [63]. The first Saturn V flight, "
      "Apollo 4 on 9 November 1967, succeeded [17]. Apollo 6, on 4 April 1968, experienced pogo oscillations, the early shutdown of "
      "two second-stage engines, and a failed third-stage restart [17]. Engineers detuned the F-1 engines, charged the "
      "liquid-oxygen prevalve cavities with helium, and redesigned the J-2 augmented spark igniter lines, and on 27 April 1968 the "
      "third Saturn V was approved for a crew [17]. Read through the model, all-up testing concentrated maturity evidence in few "
      "expensive integrated tests, Apollo 6 exposed couplings visible only at integrated level, and diagnosis and ground-verified "
      "fixes before the next integrated test correspond to what steps 4 and 5 of Table 6 would direct; the record does not give "
      "the knowledge state that the model would assign."),

("H1", "4. F-1 combustion instability"),
("P", "Between October 1962 and September 1966, Project First, the F-1 combustion-stability program, conducted about 2000 of "
      "the roughly 3200 full-scale F-1 tests, combining fourteen injection patterns with fifteen baffle configurations and rating "
      "stability with bombs producing perturbations of about 100% of chamber pressure; full-scale test data were described as the "
      "most reliable source of design information [15]. The flight injector accumulated 703 engine tests [15]. The bomb test is a "
      "controlled-perturbation test on full-scale combustion hardware: a local perturbation whose amplitude was about that of the "
      "mean chamber pressure, after which the oscillation either decays or persists, so stability becomes observable. It is not a "
      "separate-effects test in the usual sense, because the coupled combustion phenomena are all present; whether it counts as E "
      "or T in the model depends on the claim assessed (stability margin of an injector design, or the qualified engine). Read "
      "through the model, the program sits in the weak-simulation regime of Section V-C, where discriminating physical tests "
      "replace analysis as the source of knowledge. The large number of tests documents a high experimental effort; without "
      "comparable alternatives and costs, it does not show that another sequence would have been cheaper."),

("H1", "5. Falcon 1"),
("P", "Falcon 1 failed on its first three flights, in March 2006 (fuel leak and fire), March 2007 (second-stage control lost to "
      "propellant slosh), and August 2008 (first stage recontacted the second stage after separation), and reached orbit on "
      "28 September 2008 [16], [64]. Each failure revealed a different mode and was followed by diagnosis and correction. Flights "
      "cost about 8 million US dollars against roughly 100 million for development [16]. Read through the model, this is "
      "test-first development with integrated tests that were affordable relative to the program, and the program reached orbit "
      "after three failures. The record does not show how much each flight "
      "taught, and whether ground tests could have exposed these modes at lower cost is not established; Section VI-B of the "
      "article discusses why the launch-vehicle estimates support strong learning in flight only conditionally."),

("H1", "6. X-51A and HTV-2"),
("P", "The X-51A scramjet demonstrator flew four times. In May 2010 it flew for more than three minutes at Mach 4.88 under "
      "scramjet power [65]; in June 2011 the inlet unstarted during the transition to JP-7 fuel and restart attempts failed [66]; "
      "in August 2012 a cruiser control-fin fault caused the loss of the vehicle before scramjet ignition [65]; and on 1 May 2013 "
      "it flew for more than six minutes, three and a half under scramjet power, and reached Mach 5.1 [67]. The DARPA HTV-2 glider was lost "
      "on both of its flights. After the April 2010 flight, the engineering review board found higher-than-predicted adverse yaw "
      "coupled into roll that exceeded the available roll control [68]. That aerodynamic knowledge was incorporated into the August "
      "2011 flight, which achieved controlled Mach 20 flight before unexpected aeroshell degradation caused upsets of increasing "
      "severity [68], [69]. The program manager stated that extrapolating from known flight regimes and relying solely on advanced "
      "thermal modeling and ground testing could not predict the Mach 20 environment [69]."),
("P", "Read through the model, the HTV-2 statement corresponds to a low data-limited ceiling K̄: analysis and ground testing "
      "could not provide the configuration-specific predictions the flights needed, so each flight was both a demonstration and "
      "the main source of new information, and each failure exposed a different subsystem (propulsion, control surfaces, "
      "aerodynamic coupling, thermal protection). The X-51A reached a successful flight after two failures; the HTV-2 shows "
      "learning between flights without a documented terminal success."),

("H1", "7. Cross-case reading"),
("P", "Table S2-5 (Table 9 of the article) summarizes the six cases; its third column is an interpretation through the model, "
      "not a measured state. In the five flight programs, failures followed integrated tests that exposed configuration-specific "
      "or integrated-level gaps (X-43A flight 1, Apollo 6, Falcon 1 flights 1–3, X-51A flights 2–3, both HTV-2 flights). In four "
      "of them (X-43A, Saturn V, Falcon 1, X-51A) a successful flight followed diagnosis, separate-effects testing, or analysis; "
      "the HTV-2 shows that learning after failures does not guarantee a final demonstration. In the F-1 program, the absence of "
      "predictive capability led directly to controlled-perturbation tests on full-scale hardware. The cases "
      "span the scenarios of Section V-C: costly, high-consequence integrated tests (X-43A, Saturn V), weak simulation (F-1, HTV-2, "
      "X-51A), and affordable flights (Falcon 1)."),
("TABLE", ("Table S2-5. Documented sequences read through the model (Table 9 of the article). The third column is an interpretation, not a measured KMC state.",
           ["Program", "Documented sequence", "Interpretive KMC reading", "Model regime"],
           [["X-43A (2001–2004)", "Flight 1 lost; models, wind-tunnel and actuator tests revised; flights 2–3 successful [14], [55]–[59]", "Flight on heritage models → failure; tests and revised models before the next flight (restored headroom)", "Costly, risky tests"],
            ["Saturn V (1963–1968)", "All-up decision; Apollo 4 success; Apollo 6 pogo and J-2 failures; fixes; crewed Apollo 8 [17], [63]", "Few integrated tests; integrated-level failure; diagnosis and verified fixes before the next flight", "Costly, high-consequence tests"],
            ["F-1 engine (1962–1966)", "About 2000 full-scale tests with bomb rating; 703 tests of flight injector [15]", "Analysis could not predict stability; controlled-perturbation tests on full-scale hardware (low data-limited ceiling)", "Weak simulation"],
            ["Falcon 1 (2006–2008)", "Three failures with distinct modes; success on flight 4 [16], [64]", "Affordable flights at low maturity; a different failure mode each time (test-first learning)", "Cheap tests relative to program"],
            ["X-51A (2010–2013)", "Partial success; inlet unstart; fin fault; success on flight 4 [65]–[67]", "Flights were demonstrations and the main source of new information; failures on different subsystems", "Costly tests, low data-limited ceiling"],
            ["HTV-2 (2010–2011)", "Two losses: yaw–roll coupling, then aeroshell degradation [68], [69]", "Program stated that ground tests and modeling could not predict the flight environment (low data-limited ceiling); learning without a terminal success", "Costly, high-consequence tests; weak simulation"]],
           [1.2, 2.1, 2.2, 1.3])),

("H1", "8. Limitations"),
("P", "The cases were selected purposively and read with the outcomes known; the coding was not blinded or preregistered. Public documents omit export-controlled material, contemporary decision records, and activity-level "
      "costs, so no numerical K or M, no cost of alternatives, and no earlier detection or avoided loss can be established. The "
      "readings show that the mechanisms of the model are visible in documented decisions; they do not calibrate the model, test "
      "its coefficients, or show that KMC would have produced better decisions than the reviews the programs carried out. "
      "Model terms used in the readings (headroom, data-limited ceiling, K_need) are interpretations; the observable content of "
      "each case is the documented sequence of tests, failures, and corrections."),
]

# ---------------------------------------------------------------- rendering (python-docx)
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT, INK, BLUE = "Arial", RGBColor(0x1F, 0x2A, 0x37), RGBColor(0x1D, 0x5B, 0x8F)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcolor); tcPr.append(s)


def run(p, text, size=10, bold=False, italic=False, color=INK):
    """Add text; 'K_need' is typeset as K with subscript 'need'."""
    parts = text.split("K_need")
    for i, part in enumerate(parts):
        if i:
            for t, sub in (("K", False), ("need", True)):
                r = p.add_run(t); r.font.name = FONT; r.font.size = Pt(size); r.bold = bold
                r.italic = True; r.font.subscript = sub; r.font.color.rgb = color
        r = p.add_run(part); r.font.name = FONT; r.font.size = Pt(size); r.bold = bold; r.italic = italic
        r.font.color.rgb = color; r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return r


def para(doc, text, size=10, bold=False, italic=False, color=INK, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, before=0):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.15
    run(p, text, size, bold, italic, color)
    return p


def build(out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.9))
    para(doc, TITLE, 15, bold=True, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
    para(doc, SUBTITLE, 9, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=8)
    para(doc, NOTE, 9, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=10)
    width = sec.page_width - sec.left_margin - sec.right_margin
    for kind, c in BODY:
        if kind == "H1":
            para(doc, c, 12, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=4)
        elif kind == "H2":
            para(doc, c, 10.5, bold=True, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=3)
        elif kind == "P":
            para(doc, c)
        elif kind == "FIG":
            fn, cap = c
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(fn, width=Inches(5.6))
            para(doc, cap, 8.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=10)
        elif kind == "TABLE":
            cap, head, rows, widths = c
            para(doc, cap, 8.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=3)
            t = doc.add_table(rows=1, cols=len(head)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            tot = sum(widths)
            for j, h in enumerate(head):
                cell = t.rows[0].cells[j]; cell.text = ""; shade(cell, "DCE8F3")
                run(cell.paragraphs[0], h, 8, bold=True)
            for r in rows:
                cells = t.add_row().cells
                for j, v in enumerate(r):
                    cells[j].text = ""; run(cells[j].paragraphs[0], v, 8)
            for row in t.rows:
                for j, cell in enumerate(row.cells):
                    cell.width = int(width * widths[j] / tot)
            para(doc, "", 4, after=4)
    para(doc, "References (numbering of the article)", 12, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=4)
    for k in sorted(REFS):
        p = para(doc, "", 8.5, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
        run(p, f"[{k}] ", 8.5, bold=True); run(p, REFS[k], 8.5)
    cp = doc.core_properties
    cp.author = "Pedro K. de Albuquerque"; cp.last_modified_by = "Pedro K. de Albuquerque"
    cp.title = "KMC Supplementary Material S2: Documented Programs"; cp.comments = ""
    doc.save(out)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "KMC_Supplementary_S2_Documented_Programs.docx")
